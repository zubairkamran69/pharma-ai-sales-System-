import os
import re
import json
import urllib.request
import urllib.error
import io

from typing import List, Dict, Any

from dotenv import load_dotenv

PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
load_dotenv(os.path.join(PROJECT_ROOT, '.env'), override=False)


try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None


try:
    from docx import Document
except Exception:
    Document = None


try:
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
except Exception:
    TfidfVectorizer = None
    cosine_similarity = None


class MedicineRAG:
    """
    Lightweight local RAG for the demo.
    Uses TF-IDF retrieval so it runs without cloud infrastructure.
    """

    def __init__(self, db):
        self.db = db

    def extract_text(self, path: str) -> str:
        ext = os.path.splitext(path)[1].lower()

        if ext == '.pdf':
            if not PdfReader:
                raise RuntimeError(
                    'PDF parser unavailable. Install pypdf.'
                )

            reader = PdfReader(path)

            return '\n'.join(
                (p.extract_text() or '')
                for p in reader.pages
            )

        if ext == '.docx':
            if not Document:
                raise RuntimeError(
                    'DOCX parser unavailable. Install python-docx.'
                )

            doc = Document(path)

            return '\n'.join(
                p.text
                for p in doc.paragraphs
            )

        with open(
            path,
            'r',
            encoding='utf-8',
            errors='ignore'
        ) as f:
            return f.read()

    def extract_text_from_bytes(self, filename: str, content: bytes) -> str:
        ext = os.path.splitext(filename)[1].lower()
        source = io.BytesIO(content)
        if ext == '.pdf':
            if not PdfReader:
                raise RuntimeError('PDF parser unavailable. Install pypdf.')
            return '\n'.join((page.extract_text() or '') for page in PdfReader(source).pages)
        if ext == '.docx':
            if not Document:
                raise RuntimeError('DOCX parser unavailable. Install python-docx.')
            return '\n'.join(paragraph.text for paragraph in Document(source).paragraphs)
        return content.decode('utf-8', errors='ignore')

    def clean(self, text: str) -> str:
        text = text.replace('\r', '\n')

        text = re.sub(
            r'[ \t]+',
            ' ',
            text
        )

        text = re.sub(
            r'\n{3,}',
            '\n\n',
            text
        )

        return text.strip()

    def chunk(
        self,
        text: str,
        size: int = 1100,
        overlap: int = 180
    ) -> List[str]:

        text = self.clean(text)

        if not text:
            return []

        paragraphs = [
            p.strip()
            for p in re.split(
                r'\n\s*\n',
                text
            )
            if p.strip()
        ]

        chunks = []
        cur = ''

        for p in paragraphs:

            if len(cur) + len(p) + 2 <= size:
                cur = (
                    cur
                    + '\n\n'
                    + p
                ).strip()

            else:

                if cur:
                    chunks.append(cur)

                tail = (
                    cur[-overlap:]
                    if cur
                    else ''
                )

                cur = (
                    tail
                    + '\n\n'
                    + p
                ).strip()

        if cur:
            chunks.append(cur)

        return chunks

    def ingest(
        self,
        medicine_id: int,
        document_id: int,
        path: str | None,
        filename: str,
        content: bytes | None = None
    ) -> Dict[str, Any]:

        text = self.extract_text_from_bytes(filename, content) if content is not None else self.extract_text(path or '')

        chunks = self.chunk(text)

        if not chunks:
            raise ValueError(
                'No readable text found in the medicine document.'
            )

        self.db.execute(
            'DELETE FROM knowledge_chunks WHERE document_id=?',
            (document_id,)
        )

        for i, ch in enumerate(chunks, 1):

            self.db.execute(
                '''
                INSERT INTO knowledge_chunks(
                    document_id,
                    medicine_id,
                    chunk_no,
                    content,
                    metadata
                )
                VALUES(?,?,?,?,?)
                ''',
                (
                    document_id,
                    medicine_id,
                    i,
                    ch,
                    json.dumps({
                        'filename': filename,
                        'chunk': i
                    })
                )
            )

        self.db.execute(
            '''
            UPDATE documents
            SET status=?, chunk_count=?
            WHERE id=?
            ''',
            (
                'Indexed',
                len(chunks),
                document_id
            )
        )

        self.db.commit()

        return {
            'document_id': document_id,
            'chunks': len(chunks),
            'characters': len(text),
            'status': 'Indexed'
        }

    def retrieve(
        self,
        query: str,
        medicine_id: int | None = None,
        top_k: int = 5,
        owner_user_id: int | None = None
    ) -> List[Dict[str, Any]]:

        q = query.strip()

        if not q:
            return []

        if medicine_id:

            if owner_user_id:

                rows = self.db.execute(
                    '''
                    SELECT
                        kc.*,
                        d.filename,
                        m.name medicine_name,
                        m.company

                    FROM knowledge_chunks kc

                    JOIN documents d
                        ON d.id = kc.document_id

                    LEFT JOIN medicines m
                        ON m.id = kc.medicine_id

                    WHERE
                        kc.medicine_id = ?
                        AND d.owner_user_id = ?

                    ORDER BY kc.id
                    ''',
                    (
                        medicine_id,
                        owner_user_id
                    )
                ).fetchall()

            else:

                rows = self.db.execute(
                    '''
                    SELECT
                        kc.*,
                        d.filename,
                        m.name medicine_name,
                        m.company

                    FROM knowledge_chunks kc

                    JOIN documents d
                        ON d.id = kc.document_id

                    LEFT JOIN medicines m
                        ON m.id = kc.medicine_id

                    WHERE kc.medicine_id = ?

                    ORDER BY kc.id
                    ''',
                    (medicine_id,)
                ).fetchall()

        else:

            if owner_user_id:

                rows = self.db.execute(
                    '''
                    SELECT
                        kc.*,
                        d.filename,
                        m.name medicine_name,
                        m.company

                    FROM knowledge_chunks kc

                    JOIN documents d
                        ON d.id = kc.document_id

                    LEFT JOIN medicines m
                        ON m.id = kc.medicine_id

                    WHERE d.owner_user_id = ?

                    ORDER BY kc.id
                    ''',
                    (owner_user_id,)
                ).fetchall()

            else:

                rows = self.db.execute(
                    '''
                    SELECT
                        kc.*,
                        d.filename,
                        m.name medicine_name,
                        m.company

                    FROM knowledge_chunks kc

                    JOIN documents d
                        ON d.id = kc.document_id

                    LEFT JOIN medicines m
                        ON m.id = kc.medicine_id

                    ORDER BY kc.id
                    '''
                ).fetchall()

        rows = [
            dict(r)
            for r in rows
        ]

        if not rows:
            return []

        if (
            TfidfVectorizer
            and cosine_similarity
        ):

            docs = [
                r['content']
                for r in rows
            ]

            vec = TfidfVectorizer(
                stop_words='english',
                ngram_range=(1, 2),
                max_features=9000
            )

            # Fit on the documents ONLY. Words that exist only in the
            # query (Roman Urdu filler such as "ki/kya/hai", or the name
            # of a medicine that is not in the knowledge base) are then
            # ignored instead of diluting the similarity score.
            doc_mat = vec.fit_transform(docs)

            scores = cosine_similarity(
                vec.transform([q]),
                doc_mat
            ).ravel()

        else:

            terms = set(
                re.findall(
                    r'\w+',
                    q.lower()
                )
            )

            scores = [
                sum(
                    t in r['content'].lower()
                    for t in terms
                ) / max(
                    1,
                    len(terms)
                )
                for r in rows
            ]

        ranked = sorted(
            zip(rows, scores),
            key=lambda x: x[1],
            reverse=True
        )[:top_k]

        out = []

        for r, s in ranked:

            if s > 0:

                r['score'] = round(
                    float(s),
                    4
                )

                out.append(r)

        return out


class LLMClient:
    """Configurable hosted/local LLM client with an explicit extractive demo mode."""

    def __init__(self):

        self.provider = os.getenv(
            'LLM_PROVIDER',
            'groq'
        ).lower()

        self.api_key = (os.getenv('LLM_API_KEY') or os.getenv('GROQ_API_KEY') or os.getenv('OPENAI_API_KEY') or '').strip()
        model_defaults = {
            'groq': os.getenv('GROQ_MODEL', 'openai/gpt-oss-20b'),
            'openai': os.getenv('OPENAI_MODEL', 'gpt-4o-mini'),
            'openai-compatible': os.getenv('OPENAI_MODEL', 'gpt-4o-mini'),
            'ollama': os.getenv('OLLAMA_MODEL', 'llama3.2:3b')
        }
        self.model = os.getenv('LLM_MODEL') or model_defaults.get(self.provider, 'openai/gpt-oss-20b')
        default_url = {
            'groq': 'https://api.groq.com/openai/v1/chat/completions',
            'openai': 'https://api.openai.com/v1/chat/completions',
            'openai-compatible': 'https://api.openai.com/v1/chat/completions',
            'ollama': 'http://localhost:11434/api/chat'
        }.get(self.provider, '')
        provider_url = os.getenv('OLLAMA_BASE_URL') if self.provider == 'ollama' else os.getenv('OPENAI_BASE_URL')
        self.base_url = os.getenv('LLM_BASE_URL') or provider_url or default_url

        try:

            from groq import Groq

            self.client = (
                Groq(
                    api_key=self.api_key
                )
                if self.api_key
                else None
            )

        except Exception:

            self.client = None

    def generate(
        self,
        system: str,
        user: str,
        context: List[Dict[str, Any]],
        strict: bool = False
    ) -> Dict[str, Any]:
        """
        strict=True  -> medical / medicine question: the model may ONLY use
                        the retrieved sources, never general knowledge.
        strict=False -> general (non-medical) conversation.
        """

        if self.provider == 'demo':
            answer = self._fallback(user, context)
            answer['provider'] = 'demo-extractive'
            return answer
        if self.provider in ('groq', 'openai', 'openai-compatible', 'ollama'):
            return self._openai_compatible(system, user, context, strict)
        raise RuntimeError('Unsupported LLM_PROVIDER. Use groq, openai, openai-compatible, ollama, or demo.')

    def _openai_compatible(
        self,
        system: str,
        user: str,
        context: List[Dict[str, Any]],
        strict: bool = False
    ) -> Dict[str, Any]:
        if self.provider != 'ollama' and not self.api_key:
            raise RuntimeError('The selected hosted LLM provider requires LLM_API_KEY.')

        context_text = self._context(context)
        if strict:
            user_message = (
                f'User question:\n{user}\n\nApproved sources:\n{context_text}\n\n'
                'Answer only from these sources, state when evidence is missing, and cite source names.'
            )
        elif context:
            user_message = (
                f'User question:\n{user}\n\nRetrieved approved sources:\n{context_text}\n\n'
                'Use sources for product-specific facts. Do not invent unsupported facts; cite source names.'
            )
        else:
            user_message = (
                f'User question:\n{user}\n\nNo approved product-specific source was retrieved. '
                'For product facts, state that information is unavailable instead of guessing.'
            )

        payload = {'model': self.model, 'messages': [
            {'role': 'system', 'content': system},
            {'role': 'user', 'content': user_message}
        ]}
        if self.provider == 'ollama':
            payload['stream'] = False
        request_headers = {'Content-Type': 'application/json'}
        if self.provider != 'ollama':
            request_headers['Authorization'] = f'Bearer {self.api_key}'
        request = urllib.request.Request(
            self.base_url,
            data=json.dumps(payload).encode('utf-8'),
            headers=request_headers,
            method='POST'
        )
        try:
            with urllib.request.urlopen(request, timeout=45) as response:
                result = json.loads(response.read().decode('utf-8'))
        except Exception as exc:
            raise RuntimeError('The configured LLM provider request failed.') from exc
        try:
            text = result.get('message', {}).get('content') if self.provider == 'ollama' else result['choices'][0]['message']['content']
        except (KeyError, IndexError, TypeError) as exc:
            raise RuntimeError('The configured LLM provider returned an invalid response.') from exc
        if not text:
            raise RuntimeError('The configured LLM provider returned an empty response.')
        return {'text': text, 'provider': f'{self.provider}:{self.model}'}

    def _context(
        self,
        context: List[Dict[str, Any]]
    ) -> str:

        if not context:
            return ''

        return '\n\n'.join(
            f"[Source {i} | medicine: "
            f"{c.get('medicine_name') or 'company document'} | "
            f"file: {c.get('filename', '')}]\n{c['content']}"
            for i, c in enumerate(
                context,
                1
            )
        )

    def _groq(
        self,
        system: str,
        user: str,
        context: List[Dict[str, Any]],
        strict: bool = False
    ) -> Dict[str, Any]:

        try:

            if strict:

                user_message = (
                    f"User question:\n{user}\n\n"

                    "Approved knowledge-base sources:\n"

                    f"{self._context(context)}\n\n"

                    "Instructions:\n"

                    "- Answer ONLY from the sources above. Do NOT use "
                    "general or outside knowledge.\n"

                    "- If the sources do not contain the answer, say "
                    "that the approved knowledge base does not contain "
                    "it.\n"

                    "- Cite sources like [Source 1].\n"

                    "- Reply in the user's language.\n"
                )

            elif context:

                user_message = (
                    f"User question:\n{user}\n\n"

                    "Retrieved approved knowledge:\n"

                    f"{self._context(context)}\n\n"

                    "Instructions:\n"

                    "- Use the retrieved knowledge for "
                    "product/company-specific facts.\n"

                    "- Do not invent product-specific facts "
                    "that are not supported by the retrieved sources.\n"

                    "- You may explain, summarize, organize, "
                    "or compare supported information.\n"

                    "- For general questions, answer naturally "
                    "using your general knowledge.\n"
                )

            else:

                user_message = (
                    f"User question:\n{user}\n\n"

                    "No relevant approved product-document "
                    "information was retrieved.\n\n"

                    "For a general question, answer normally "
                    "using your general knowledge.\n\n"

                    "For a product-specific pharmaceutical question "
                    "where the required information is not available "
                    "in the approved knowledge base, clearly say that "
                    "the information is unavailable instead of "
                    "guessing or inventing it."
                )

            response = (
                self.client
                .chat
                .completions
                .create(
                    model=self.model,

                    messages=[
                        {
                            'role': 'system',
                            'content': system
                        },
                        {
                            'role': 'user',
                            'content': user_message
                        }
                    ],

                    temperature=0.2,

                    max_tokens=1000
                )
            )

            answer = (
                response
                .choices[0]
                .message
                .content
            )

            return {
                'text': answer,
                'provider': (
                    'groq:'
                    + self.model
                )
            }

        except Exception as e:
            raise RuntimeError('The configured Groq request failed.') from e

    def _fallback(
        self,
        user: str,
        context: List[Dict[str, Any]]
    ) -> Dict[str, Any]:

        if not context:

            return {
                'text': (
                    "I don't have approved "
                    "product-specific information "
                    "for that question. Please upload "
                    "or index the relevant document."
                ),
                'provider': 'fallback'
            }

        q = user.lower()

        sentences = []

        for c in context:

            sentences.extend(
                re.split(
                    r'(?<=[.!?])\s+',
                    c['content'].replace(
                        '\n',
                        ' '
                    )
                )
            )

        sentences = [
            re.sub(
                r'\s+',
                ' ',
                x
            ).strip()

            for x in sentences

            if len(x.strip()) > 20
        ]

        terms = set(
            re.findall(
                r'\w+',
                q
            )
        )

        scored = []

        for sentence in sentences:

            sentence_terms = set(
                re.findall(
                    r'\w+',
                    sentence.lower()
                )
            )

            overlap = len(
                sentence_terms & terms
            )

            if overlap:

                scored.append(
                    (
                        overlap,
                        sentence
                    )
                )

        scored.sort(
            reverse=True
        )

        focus = [
            sentence
            for _, sentence
            in scored[:3]
        ]

        if focus:

            answer = ' '.join(
                focus
            )

        else:

            answer = (
                'I found approved information, '
                'but it does not contain enough '
                'evidence to answer the question precisely.'
            )

        return {
            'text': answer,
            'provider': 'fallback'
        }